"""Convert the team's BO_/SG_ workbook layout without a desktop file picker."""

import argparse
import math
import re
from pathlib import Path

import cantools
from openpyxl import load_workbook


class WorkbookError(ValueError):
    """A definition is invalid; never silently repair a CAN interface."""


def identifier(value):
    name = str(value or "").strip().replace(" ", "_")
    if not re.fullmatch(r"[A-Za-z_][A-Za-z_0-9]*", name):
        raise WorkbookError(f"Invalid or missing identifier: {value!r}")
    return name


def number(row, column, default=None, integer=False):
    value = row.get(column)
    if value is None or value == "":
        if default is None:
            raise WorkbookError(f"Missing {column}")
        value = default
    try:
        result = float(value)
    except (ValueError, TypeError):
        raise WorkbookError(f"Invalid {column}: {value!r}") from None
    if not math.isfinite(result) or (integer and result != int(result)):
        raise WorkbookError(f"Invalid {column}: {value!r}")
    return int(result) if integer else result


def receivers(value, fallback="Vector__XXX"):
    if value is None or not str(value).strip():
        return fallback
    return ",".join(identifier(part) for part in str(value).split(","))


def signal_line(row, fallback):
    name = identifier(row.get("Signal Name"))
    start = number(row, "Start Bit", integer=True)
    length = number(row, "Length", integer=True)
    if start < 0 or not 1 <= length <= 64:
        raise WorkbookError("Start Bit must be nonnegative; Length must be 1..64")
    order = str(row.get("Byte Order", "")).strip().lower()
    kind = str(row.get("Data Type", "")).strip().lower()
    if order not in ("intel", "motorola") or kind not in ("signed", "unsigned", "bool"):
        raise WorkbookError("Expected Intel/Motorola Byte Order and Signed/Unsigned/Bool Data Type")
    factor = number(row, "Factor", 1)
    offset = number(row, "Offset", 0)
    minimum = number(row, "Minimum", 0)
    maximum = number(row, "Maximum", 0)
    if factor == 0 or minimum > maximum:
        raise WorkbookError("Factor must be nonzero and Minimum must not exceed Maximum")
    unit = str(row.get("Unit") or "").strip()
    if any(char in unit for char in ('"', '\n', '\r', '\\')):
        raise WorkbookError("Unit contains characters unsupported by the DBC header generator")
    mux = str(row.get("Multiplex") if row.get("Multiplex") is not None else "").strip()
    if mux.isdigit():
        mux = "m" + mux
    if mux and not re.fullmatch(r"M|m[0-9]+", mux):
        raise WorkbookError(f"Invalid Multiplex value: {mux!r}")
    mux = " " + mux if mux else ""
    byte_order = 0 if order == "motorola" else 1
    sign = "-" if kind == "signed" else "+"
    return name, (f" SG_ {name}{mux} : {start}|{length}@{byte_order}{sign} "
                  f"({factor}, {offset}) [{minimum}|{maximum}] "
                  f'"{unit}" {receivers(row.get("Receivers"), fallback)}')


def value_table(value, frame_id, signal):
    if value is None or not str(value).strip():
        return None
    entries, seen = [], set()
    for token in str(value).split(","):
        try:
            raw, label = token.split("=", 1)
            raw = raw.strip()
            key = int(raw, 16) if raw.lower().startswith(("0x", "-0x")) else int(raw)
        except ValueError:
            raise WorkbookError(f"Invalid Value Table entry: {token!r}") from None
        label = label.strip()
        if key in seen or not label or any(c in label for c in ('"', '\\', '\n', '\r')):
            raise WorkbookError(f"Duplicate value or invalid label in Value Table: {token!r}")
        seen.add(key)
        entries.append(f'{key} "{label}"')
    return f"VAL_ {frame_id} {signal} {' '.join(entries)} ;"


def convert(workbook_path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    nodes, messages, attributes, mappings = [], [], [], []
    ids, names = set(), set()
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headings = next(rows, ())
            if "Type" not in headings:
                continue  # Diagnostics/instruction tabs, including CAN Plan.
            headings = [str(value).strip() if value is not None else "" for value in headings]
            required = {"Type", "Msg Name", "ID (Hex)", "DLC", "Signal Name", "Start Bit",
                        "Length", "Byte Order", "Data Type"}
            if not required.issubset(headings) or len(set(filter(None, headings))) != len(list(filter(None, headings))):
                raise WorkbookError(f"{sheet.title}: missing or duplicate CAN column headings")
            node = identifier(sheet.title)
            if node in nodes:
                raise WorkbookError(f"Duplicate normalized board name: {node}")
            nodes.append(node)
            current_id, inherited, signals = None, "Vector__XXX", set()
            for index, values in enumerate(rows, 2):
                row = dict(zip(headings, values))
                if not any(value is not None and str(value).strip() for key, value in row.items() if key):
                    continue
                try:
                    if any(isinstance(v, str) and v.startswith("=") for k, v in row.items() if k):
                        raise WorkbookError("Use literal values in CAN definition cells; formulas belong in CAN Plan")
                    kind = str(row.get("Type") or "").strip().upper()
                    if kind == "BO_":
                        name = identifier(row.get("Msg Name"))
                        raw_id = row.get("ID (Hex)")
                        current_id = int(raw_id.strip(), 16) if isinstance(raw_id, str) else number(row, "ID (Hex)", integer=True)
                        dlc = number(row, "DLC", integer=True)
                        if not 0 <= current_id <= 0x7FF or not 0 <= dlc <= 8:
                            raise WorkbookError("Expected Classic CAN ID 0x000..0x7FF and DLC 0..8")
                        if current_id in ids or name in names:
                            raise WorkbookError(f"Duplicate message ID 0x{current_id:X} or name {name}")
                        ids.add(current_id)
                        names.add(name)
                        signals = set()
                        inherited = receivers(row.get("Receivers"))
                        messages.append(f"\nBO_ {current_id} {name}: {dlc} {node}")
                        frequency = row.get("Frequency (Hz)")
                        if frequency is not None:
                            try:
                                float(frequency)
                            except (TypeError, ValueError):
                                pass  # Event-driven frequencies remain descriptive text.
                            else:
                                hz = number(row, "Frequency (Hz)")
                                if hz <= 0 or not 1 <= int(1000 / hz) <= 65535:
                                    raise WorkbookError("Frequency must map to a cycle time of 1..65535 ms")
                                attributes.append(f'BA_ "GenMsgCycleTime" BO_ {current_id} {int(1000 / hz)};')
                    elif kind == "SG_":
                        if current_id is None:
                            raise WorkbookError("Signal appears before a BO_ message")
                        name, line = signal_line(row, inherited)
                        if name in signals:
                            raise WorkbookError(f"Duplicate signal {name}")
                        signals.add(name)
                        messages.append(line)
                        mapping = value_table(row.get("Value Table"), current_id, name)
                        if mapping:
                            mappings.append(mapping)
                    else:
                        raise WorkbookError(f"Expected BO_ or SG_ in Type, got {kind!r}")
                except (ValueError, TypeError) as error:
                    raise WorkbookError(f"{sheet.title}!row {index}: {error}") from error
    finally:
        workbook.close()
    if not ids:
        raise WorkbookError("Workbook contains no CAN messages")
    text = '\n'.join(['VERSION ""', '', 'NS_ :', '\tBA_', '\tBA_DEF_', '\tBA_DEF_DEF_', '\tVAL_',
                      '', 'BS_:', '', 'BU_: ' + ' '.join(nodes), *messages, '',
                      'BA_DEF_ BO_ "GenMsgCycleTime" INT 0 65535;',
                      'BA_DEF_DEF_ "GenMsgCycleTime" 0;', *attributes, *mappings, ''])
    try:
        cantools.database.load_string(text, database_format="dbc", strict=True, sort_signals=None)
    except Exception as error:
        raise WorkbookError(f"DBC validation failed (message/signal named below): {error}") from error
    return text


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("MAIN_DBC.dbc"))
    args = parser.parse_args()
    try:
        text = convert(args.workbook)
    except (WorkbookError, OSError) as error:
        parser.exit(1, f"{error}\n")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(text, encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
