"""Regression tests for wire definitions and direct-to-main publication."""

import contextlib
import io
import os
import subprocess
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

import cantools
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from download import download, download_url
from excel_to_dbc import WorkbookError, convert
from pipeline import GENERATED, PROVENANCE, next_tag, prepare, publish, release_tag, stage


HEADINGS = ["Type", "Msg Name", "ID (Hex)", "DLC", "Frequency (Hz)", "Signal Name",
            "Start Bit", "Length", "Byte Order", "Data Type", "Factor", "Minimum",
            "Maximum", "Unit", "Receivers", "Value Table", "Multiplex"]
ROOT = Path(__file__).resolve().parents[3]


def workbook(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "BMS Master"
    sheet.append(HEADINGS)
    for row in rows:
        sheet.append([row.get(key) for key in HEADINGS])
    book.create_sheet("CAN Plan").append(["Diagnostics"])
    book.save(path)
    book.close()


def message(**overrides):
    return {"Type": "BO_", "Msg Name": "BM_Status", "ID (Hex)": "0x010", "DLC": 8,
            "Frequency (Hz)": 10, "Receivers": "Telemetry, Dashboard", **overrides}


def signal(**overrides):
    return {"Type": "SG_", "Signal Name": "Voltage", "Start Bit": 7, "Length": 16,
            "Byte Order": "Motorola", "Data Type": "Unsigned", "Factor": 0.1,
            "Minimum": 0, "Maximum": 6553.5, "Unit": "V", **overrides}


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "source.xlsx"

    def test_preserves_endianness_scaling_receivers_and_cycle(self):
        workbook(self.path, [message(), signal()])
        text = convert(self.path)
        db = cantools.database.load_string(text, database_format="dbc")
        msg = db.get_message_by_name("BM_Status")
        self.assertEqual(msg.frame_id, 0x10)
        self.assertEqual(msg.cycle_time, 100)
        self.assertEqual(msg.signals[0].receivers, ["Telemetry", "Dashboard"])
        data = msg.encode({"Voltage": 123.4})
        self.assertEqual(data[:2], bytes([0x04, 0xD2]))
        self.assertAlmostEqual(msg.decode(data)["Voltage"], 123.4)
        self.assertNotIn("CAN_Plan", text)

    def test_mux_values_and_signed_signal(self):
        workbook(self.path, [message(), signal(**{"Signal Name": "Mode", "Start Bit": 0,
            "Length": 8, "Byte Order": "Intel", "Factor": 1, "Maximum": 1,
            "Multiplex": "M", "Value Table": "0=STOPPED, 1=RUNNING"}),
            signal(**{"Signal Name": "Current", "Start Bit": 8, "Byte Order": "Intel",
            "Data Type": "Signed", "Minimum": -3276.8, "Maximum": 3276.7, "Multiplex": 1})])
        msg = cantools.database.load_string(convert(self.path), database_format="dbc").messages[0]
        self.assertEqual(str(msg.signals[0].choices[1]), "RUNNING")
        self.assertAlmostEqual(msg.decode(msg.encode({"Mode": 1, "Current": -12.3}))["Current"], -12.3)

    def test_rejects_invalid_definitions_with_sheet_and_row(self):
        cases = [(signal(Factor="q"), "Factor"), (signal(Length=0), "Length"),
                 (signal(**{"Start Bit": 60}), "DBC validation"),
                 (signal(**{"Byte Order": "mystery"}), "Byte Order"),
                 (signal(Factor="=A2"), "literal values"),
                 (signal(**{"Value Table": "1=OK, 1=BAD"}), "Duplicate value")]
        for row, expected in cases:
            with self.subTest(expected=expected):
                workbook(self.path, [message(), row])
                with self.assertRaisesRegex(WorkbookError, expected):
                    convert(self.path)
        workbook(self.path, [message(), signal(Factor="q")])
        with self.assertRaisesRegex(WorkbookError, "BMS Master!row 3"):
            convert(self.path)

    def test_rejects_duplicate_ids_names_signals_and_overlaps(self):
        cases = [[message(), signal(), message()], [message(), signal(), signal()],
                 [message(), signal(), signal(**{"Signal Name": "Other"})],
                 [message(**{"ID (Hex)": "garbage"}), signal()],
                 [message(**{"ID (Hex)": "0x800"}), signal()], [signal()]]
        for rows in cases:
            with self.subTest(rows=rows):
                workbook(self.path, rows)
                with self.assertRaises(WorkbookError):
                    convert(self.path)

    def test_formatting_only_edit_does_not_stage_or_bump(self):
        workbook(self.path, [message(), signal()])
        artifacts = Path(self.temp.name) / "artifacts"
        checkout = Path(self.temp.name) / "checkout"
        prepare(self.path, artifacts)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertTrue(stage(artifacts, checkout))
        provenance = (checkout / PROVENANCE).read_bytes()
        book = load_workbook(self.path)
        book.active.column_dimensions["A"].width = 40
        book.save(self.path)
        book.close()
        prepare(self.path, artifacts)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertFalse(stage(artifacts, checkout))
        self.assertEqual((checkout / PROVENANCE).read_bytes(), provenance)

    def test_failed_generation_leaves_existing_outputs_untouched(self):
        workbook(self.path, [message(), signal(Factor="q")])
        artifacts = Path(self.temp.name) / "artifacts"
        artifacts.mkdir()
        (artifacts / "MAIN_DBC.dbc").write_text("existing")
        with self.assertRaises(WorkbookError):
            prepare(self.path, artifacts)
        self.assertEqual((artifacts / "MAIN_DBC.dbc").read_text(), "existing")


class DownloadTests(unittest.TestCase):
    def test_replaces_download_parameter(self):
        url = download_url("https://team.sharepoint.com/:x:/s/team/token?e=abc&download=0#fragment")
        self.assertEqual(url, "https://team.sharepoint.com/:x:/s/team/token?e=abc&download=1")
        for bad in ("http://team.sharepoint.com/file", "https://sharepoint.com.evil.test/file"):
            with self.assertRaises(ValueError):
                download_url(bad)

    def test_guest_cookie_redirect_and_rejects_login_html(self):
        book = Workbook()
        data = io.BytesIO()
        book.save(data)
        book.close()
        content = data.getvalue()

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                if self.path == "/share":
                    self.send_response(302)
                    self.send_header("Set-Cookie", "guest=ok; Path=/")
                    self.send_header("Location", "/file")
                    self.end_headers()
                else:
                    self.send_response(200)
                    self.end_headers()
                    self.wfile.write(content if self.headers.get("Cookie") == "guest=ok" else b"<html>Sign in</html>")

            def log_message(self, *args):
                pass

        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / "source.xlsx"
                base = f"http://127.0.0.1:{server.server_port}"
                with patch("download.download_url", return_value=base + "/share"):
                    self.assertIn("workbook_sha256", download("unused", path))
                self.assertEqual(path.read_bytes(), content)
                with patch("download.download_url", return_value=base + "/file"):
                    with self.assertRaisesRegex(ValueError, "sign-in/error page"):
                        download("unused", path)
                self.assertEqual(path.read_bytes(), content)
        finally:
            server.shutdown()
            server.server_close()
            thread.join()


class PublicationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.addCleanup(patch.stopall)
        patch.dict(os.environ, {"GITHUB_OUTPUT": ""}).start()
        base = Path(self.temp.name)
        self.origin, self.root = base / "origin.git", base / "checkout"
        self.run_git(base, "init", "--bare", "--initial-branch=main", str(self.origin))
        self.run_git(base, "clone", str(self.origin), str(self.root))
        self.run_git(self.root, "config", "user.name", "Test")
        self.run_git(self.root, "config", "user.email", "test@example.invalid")
        for target in [*GENERATED.values(), PROVENANCE]:
            path = self.root / target
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("before")
        self.run_git(self.root, "add", ".")
        self.run_git(self.root, "commit", "-m", "Initial")
        self.run_git(self.root, "tag", "v1.4.2")
        self.run_git(self.root, "push", "origin", "main", "--tags")
        (self.root / "MAIN_DBC.dbc").write_text("after")
        (self.root / PROVENANCE).write_text("new source")

    @staticmethod
    def run_git(root, *args):
        return subprocess.check_output(["git", "-C", str(root), *args], stderr=subprocess.STDOUT, text=True).strip()

    def test_semantic_patch_increment(self):
        self.assertEqual(next_tag(["v1.4.2", "v1.4.9", "v1.4.10", "1.3.18", "v2.0.0-rc1"]), "v1.4.11")
        self.assertEqual(next_tag(["v1.5.0", "v1.4.99"]), "v1.5.1")
        with self.assertRaises(ValueError):
            next_tag(["nightly"])

    def test_atomic_commit_and_tag_and_release_recovery(self):
        with contextlib.redirect_stdout(io.StringIO()):
            publish(self.root)
        head = self.run_git(self.origin, "rev-parse", "main")
        self.assertEqual(head, self.run_git(self.origin, "rev-parse", "v1.4.3"))
        self.assertEqual(self.run_git(self.origin, "show", "main:MAIN_DBC.dbc"), "after")
        self.assertEqual(self.run_git(self.root, "status", "--porcelain"), "")
        # A later unrelated commit must not lose the tag needed to retry a release.
        (self.root / "README.md").write_text("Later documentation")
        self.run_git(self.root, "add", "README.md")
        self.run_git(self.root, "commit", "-m", "Docs")
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            release_tag(self.root)
        self.assertEqual(stream.getvalue().strip(), "tag=v1.4.3")

    def test_refuses_unrelated_files(self):
        (self.root / "unrelated.txt").write_text("user work")
        with self.assertRaisesRegex(ValueError, "Unexpected repository changes"):
            publish(self.root)
        self.assertEqual(self.run_git(self.origin, "tag", "--list"), "v1.4.2")

    def test_refuses_main_that_advanced_while_generation_ran(self):
        old = self.run_git(self.origin, "rev-parse", "main")
        self.run_git(self.origin, "config", "user.name", "Test")
        self.run_git(self.origin, "config", "user.email", "test@example.invalid")
        new = subprocess.check_output(["git", "-C", str(self.origin), "commit-tree",
              "main^{tree}", "-p", old, "-m", "Concurrent change"], text=True).strip()
        self.run_git(self.origin, "update-ref", "refs/heads/main", new)
        with self.assertRaisesRegex(ValueError, "main advanced"):
            publish(self.root)
        self.assertEqual(self.run_git(self.origin, "tag", "--list"), "v1.4.2")

    def test_remote_tag_race_cannot_push_main_without_tag(self):
        old = self.run_git(self.origin, "rev-parse", "main")
        self.run_git(self.origin, "config", "user.name", "Test")
        self.run_git(self.origin, "config", "user.email", "test@example.invalid")
        other = subprocess.check_output(["git", "-C", str(self.origin), "commit-tree",
                 "main^{tree}", "-m", "Other tag target"], text=True).strip()
        original = subprocess.check_output

        def race(command, **kwargs):
            if command[3:5] == ["push", "--atomic"]:
                self.run_git(self.origin, "update-ref", "refs/tags/v1.4.3", other)
            return original(command, **kwargs)

        with patch("pipeline.subprocess.check_output", side_effect=race):
            with self.assertRaises(subprocess.CalledProcessError):
                publish(self.root)
        self.assertEqual(self.run_git(self.origin, "rev-parse", "main"), old)


if __name__ == "__main__":
    unittest.main()
